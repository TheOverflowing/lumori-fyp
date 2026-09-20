"""Build the formal Part-I manuscript. No model calls or application writes."""
from pathlib import Path
import csv, hashlib, json, re, sys, os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'docs/thesis'; FIG=BASE/'figures'; OUT=ROOT/'output'; OUT.mkdir(exist_ok=True)
# Use native raster drawing for the simple architecture; a vector chart is also saved.
FONT='/System/Library/Fonts/Supplemental/Arial.ttf'
BOLD='/System/Library/Fonts/Supplemental/Arial Bold.ttf'
def font(n,b=False): return ImageFont.truetype(BOLD if b else FONT,n)
def architecture():
    im=Image.new('RGB',(1800,730),'white'); d=ImageDraw.Draw(im)
    def box(x,y,w,h,title,body):
        d.rectangle((x,y,x+w,y+h),fill='#f7f8fa',outline='#737b86',width=2)
        d.text((x+w/2,y+34),title,font=font(30,True),fill='#162331',anchor='mm')
        for j,line in enumerate(body): d.text((x+w/2,y+80+j*35),line,font=font(25),fill='#273440',anchor='mm')
    def arrow(x,y,xx,yy):
        d.line((x,y,xx,yy),fill='#344c61',width=4)
        if xx>x:d.polygon([(xx,yy),(xx-15,yy-9),(xx-15,yy+9)],fill='#344c61')
        else:d.polygon([(xx,yy),(xx-9,yy-15),(xx+9,yy-15)],fill='#344c61')
    box(20,30,350,160,'Browser interface',['Authoring and practice','Owned course workspace'])
    box(460,30,420,160,'FastAPI coordinator',['Identity and ownership','Jobs and validation'])
    box(980,30,790,160,'Persistence and provenance',['SQLite + local files','Sources, vectors, asset versions, evidence'])
    arrow(370,110,460,110);arrow(880,110,980,110)
    box(20,300,350,170,'Ingestion',['Native text / MinerU','Optional figure branch'])
    box(460,300,420,170,'Retrieval',['Recursive chunks + vectors','Dense / lexical → rerank'])
    box(980,300,340,170,'Generation',['DeepSeek text model','Grounded draft'])
    box(1410,300,360,170,'Validation',['Structure + demand','Versioned result'])
    arrow(370,385,460,385);arrow(880,385,980,385);arrow(1320,385,1410,385)
    arrow(195,190,195,300);arrow(670,190,670,300)
    d.line((1585,300,1585,190),fill='#344c61',width=3)
    d.polygon([(1585,190),(1576,205),(1594,205)],fill='#344c61')
    d.text((895,570),'Provider adapters: text generation · embeddings · reranking',font=font(28),fill='#273440',anchor='mm')
    d.text((895,622),'Selected sources and generated text are recorded as distinct, versioned evidence.',font=font(26),fill='#273440',anchor='mm')
    im.save(FIG/'architecture.png')

def chart():
    # ReportLab drawing retains a vector original for publication reuse.
    from reportlab.graphics.shapes import Drawing, Rect, String, Line
    from reportlab.graphics import renderPDF
    from reportlab.lib.colors import HexColor
    D=Drawing(600,220)
    labels=['Baseline','Instruction only','Selected API 8B']
    vals=[.8868,.9484,.9524]
    colors=['#a0aab4','#607e98','#183d5b']
    left=145;right=550;y0=45;top=190
    for t in [0,.25,.5,.75,1.0]:
        x=left+(right-left)*t
        D.add(Line(x,y0,x,top,strokeColor=HexColor('#dddddd'),strokeWidth=.5))
        D.add(String(x,25,f'{t:.2f}',fontName='Helvetica',fontSize=10,textAnchor='middle'))
    for i,(label,val,col) in enumerate(zip(labels,vals,colors)):
        y=160-i*45
        D.add(String(left-10,y+9,label,fontName='Helvetica',fontSize=11,textAnchor='end'))
        D.add(Rect(left,y,(right-left)*val,27,fillColor=HexColor(col),strokeColor=None))
        D.add(String(left+(right-left)*val+5,y+8,f'{val:.4f}',fontName='Helvetica',fontSize=10))
    D.add(String(345,5,'Held-out MRR (n = 168 scenarios)',fontName='Helvetica',fontSize=11,textAnchor='middle'))
    renderPDF.drawToFile(D,str(FIG/'r4-results.pdf'))
    import subprocess
    poppler=Path(sys.executable).parent.parent.parent/'bin/override/pdftoppm'
    subprocess.run([str(poppler),'-png','-scale-to','1800','-singlefile',str(FIG/'r4-results.pdf'),str(FIG/'r4-results')],check=True,stdout=subprocess.DEVNULL)

# Keep fontconfig caches inside the project on restricted desktop hosts.
fc=ROOT/'tmp/thesis/fonts.conf'
if fc.exists(): os.environ['FONTCONFIG_FILE']=str(fc)
architecture();chart()
rows=list(csv.DictReader((ROOT/'docs/ca1/literature/reference_index.csv').open()))
refs=[r['citation']+' '+r['url'] for r in rows]
refs += [
'R. Qu, R. Tu, and F. S. Bao. Is Semantic Chunking Worth the Computational Cost? Findings of NAACL, 2025, pp. 2155–2177. https://aclanthology.org/2025.findings-naacl.114/',
'M. Günther, I. Mohr, D. J. Williams, B. Wang, and H. Xiao. Late Chunking: Contextual Chunk Embeddings Using Long-Context Embedding Models. arXiv:2409.04701, 2024; version 3, 2025. Preprint. https://arxiv.org/abs/2409.04701',
'Y. Zhang et al. Qwen3 Embedding: Advancing Text Embedding and Reranking Through Foundation Models. arXiv:2506.05176, 2025. Preprint. https://arxiv.org/abs/2506.05176',
'P. Lewis et al. MLQA: Evaluating Cross-lingual Extractive Question Answering. ACL, 2020, pp. 7315–7330. https://aclanthology.org/2020.acl-main.653/',
'X. Xie et al. T2Ranking: A large-scale Chinese Benchmark for Passage Ranking. SIGIR, 2023. Author version: arXiv:2304.03679. https://arxiv.org/abs/2304.03679'
]
# Optional inspected page-map makes the contents a stable, accessible static table.
mapfile=BASE/'page_map.json'; pagemap=json.loads(mapfile.read_text()) if mapfile.exists() else {}
contents=[('1  Introduction','4'),('2  Literature Review','5'),('3  Requirements and System Design','7'),('4  Implementation','9'),('5  Evaluation Methodology','12'),('6  Results and Discussion','14'),('7  Limitations and Next Research Stage','19'),('References','20'),('Appendix A  AI-Assisted Content Audit','22'),('Appendix B  Evidence and Reproduction Index','24')]
text=(BASE/'Manuscript.md').read_text()
text=text.replace('{{CONTENTS}}','\n'.join(f'{h} | {pagemap.get(h,p)}' for h,p in contents))
text=text.replace('{{REFERENCES_1}}','\n\n'.join(f'[{i+1}] {v}' for i,v in enumerate(refs[:10])))
text=text.replace('{{REFERENCES_2}}','\n\n'.join(f'[{i+11}] {v}' for i,v in enumerate(refs[10:])))
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.3)
sec.header_distance=Cm(1.0);sec.footer_distance=Cm(1.0)
sec.different_first_page_header_footer=True
styles=doc.styles
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[name].font.name='Times New Roman';styles[name].font.color.rgb=RGBColor(0,0,0)
    rf=styles[name].element.get_or_add_rPr().rFonts
    for attr in list(rf.attrib):
        if 'Theme' in attr: del rf.attrib[attr]
    for attr in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+attr),'Times New Roman')
normal=styles['Normal'];normal.font.size=Pt(11.5)
normal.paragraph_format.line_spacing=1.10
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for n,size,before,after in [('Heading 1',18,0,15),('Heading 2',13,12,7),('Heading 3',11.5,10,6)]:
    s=styles[n];s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after)
    s.paragraph_format.keep_with_next=True
for sty in ['Title','Subtitle']:
    pp=styles[sty].element.find(qn('w:pPr'))
    if pp is not None:
        for b in list(pp.findall(qn('w:pBdr'))): pp.remove(b)
styles['Caption'].font.size=Pt(10);styles['Caption'].font.italic=False
styles['Caption'].paragraph_format.space_after=Pt(8)
styles['Caption'].paragraph_format.line_spacing=1.05
# Running head and automatic page numbers.
hp=sec.header.paragraphs[0];hp.text='LUMORI  /  CORE TEXT GENERATION AND RETRIEVAL';hp.style='Caption'
hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor.from_string('606060')
pp=hp._p.get_or_add_pPr();border=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
for k,v in {'val':'single','sz':'4','color':'BFBFBF','space':'5'}.items():bottom.set(qn('w:'+k),v)
border.append(bottom);pp.append(border)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
# Use explicit authored page groups so related argument, table and interpretation stay together.
def para(s,style=None):
    p=doc.add_paragraph(s,style)
    if style is None:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def table(lines):
    data=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
    data=[r for r in data if not all(re.fullmatch(r'[-: ]+',c) for c in r)]
    t=doc.add_table(rows=0,cols=len(data[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths={3:[1.65,2.30,2.43],4:[2.2,1.40,1.39,1.39],5:[2.35,1.03,1.1,1.08,.82]}.get(len(data[0]),[6.38/len(data[0])]*len(data[0]))
    if data[0][0]=='ID':widths=[.42,3.37,2.59]
    if data[0][0]=='Review ID':widths=[.8,.82,4.76]
    if data[0][0]=='Step':widths=[.5,5.88]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D4D4D4');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'80' if side in ['top','bottom'] else '90');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for i,values in enumerate(data):
        row=t.add_row();trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,s) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
            # Permit long evidence paths to wrap without inserting visible punctuation.
            s=s.replace('/','/\u200b') if 'docs/' in s else s
            r=p.add_run(s);r.font.size=Pt(9.5 if len(data[0])>3 or data[0][0]=='ID' else 10)
            if i==0:
                r.bold=True;shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F0');cell._tc.get_or_add_tcPr().append(shade)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)

for page_idx,chunk in enumerate(text.split('<!-- PAGE -->')):
    start_paragraph=len(doc.paragraphs)
    if page_idx==0:
        p=doc.add_paragraph('LUMORI',style='Title');p.paragraph_format.space_before=Pt(80)
        p.runs[0].font.size=Pt(34);p.runs[0].bold=True
        p=doc.add_paragraph('Course-Grounded Educational\nContent and Assessment Generation');p.paragraph_format.space_before=Pt(22);p.paragraph_format.space_after=Pt(28)
        p.runs[0].font.size=Pt(23)
        p=doc.add_paragraph('INTERIM THESIS MANUSCRIPT · PART I');p.runs[0].font.size=Pt(11);p.runs[0].bold=True
        para('Core text generation, retrieval and difficulty control')
        p=para('AI-Powered Multimodal Content & Assessment Generation for Educational Applications');p.paragraph_format.space_before=Pt(30)
        para('Final-year project · CA1 stage · AY 2026/2027')
        para('19 September 2026  |  Version 0.2')
        p=para('Working thesis manuscript\nMultimodal output research will be incorporated in a subsequent stage.');p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_before=Pt(58);p.runs[0].font.size=Pt(10);p.runs[0].italic=True
        continue
    lines=chunk.strip().splitlines();i=0
    while i<len(lines):
        s=lines[i].strip();i+=1
        if not s:continue
        if s.startswith('|'):
            block=[s]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
            table(block);continue
        if s.startswith('#'):
            n=len(s)-len(s.lstrip('#'));para(s[n:].strip(),f'Heading {min(n,3)}');continue
        if s.startswith('!['):
            m=re.fullmatch(r'!\[(.*)\]\((.*)\)',s);p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            width=4.8 if 'difficulty-detail' in m[2] else (5.8 if 'architecture' in m[2] else 6.3)
            p.add_run().add_picture(str(BASE/m[2]),width=Inches(width))
            para(m[1],'Caption');continue
        if page_idx==2 and ' | ' in s:
            label,num=s.rsplit(' | ',1);p=doc.add_paragraph();p.paragraph_format.space_after=Pt(12)
            from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.3),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            p.add_run(label+'\t'+num);continue
        if s.startswith('Table ') or s.startswith('Algorithm '):
            p=para(s,'Caption');p.paragraph_format.keep_with_next=True;continue
        if re.match(r'^\[\d+\]',s):
            p=para(s);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_after=Pt(11);p.paragraph_format.line_spacing=1.07
            p.paragraph_format.left_indent=Cm(.7);p.paragraph_format.first_line_indent=Cm(-.7)
            for r in p.runs:r.font.size=Pt(10.5)
            continue
        para(s)
    if page_idx: doc.paragraphs[start_paragraph].paragraph_format.page_break_before=True
props=doc.core_properties;props.title='Lumori: Course-Grounded Educational Content and Assessment Generation';props.subject='Interim thesis manuscript, Part I';props.author='Lumori project';props.keywords='RAG, educational generation, difficulty control, CA1';props.comments='AI-assisted manuscript. Historical experiments are dated; no human review is asserted.'
out=OUT/'Lumori_Thesis_Part_I_v0.2.docx';doc.save(out)
manifest={'capture_date':'2026-09-19','mode':'light','source':'actual local application, http://127.0.0.1:8765/','notes':'Screenshots show current UI and historical assets; difficulty form state is not a new generation result. No application records changed.','figures':[{'file':p.name,'kind':'browser_capture' if p.name.startswith('ui-') else 'analytical_figure','sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(FIG.glob('*.png'))]}
(BASE/'figure_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
(BASE/'references.json').write_text(json.dumps([{'number':i+1,'citation':v} for i,v in enumerate(refs)],ensure_ascii=False,indent=2)+'\n')
print(out)
