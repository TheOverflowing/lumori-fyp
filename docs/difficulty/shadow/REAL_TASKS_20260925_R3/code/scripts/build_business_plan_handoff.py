#!/usr/bin/env python3
"""Build the English Lumori business-plan handoff from its Markdown source.

Only creates presentation assets and DOCX. It does not start Lumori or call any model.
"""

from __future__ import annotations

import math
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/business_plan/HANDOFF_20260923.md"
ASSETS = SOURCE.parent / "assets"
OUT = ROOT / "output/business_plan_handoff_20260923"
DOCX = OUT / "Lumori_Business_Plan_Handoff_20260923.docx"
FONTCONFIG = OUT / "fonts.conf"

NAVY = "203B50"
BLUE = "416C88"
TEAL = "167E85"
PALE = "EDF5F5"
PALE_BLUE = "F1F6FA"
AMBER = "D48A3A"
INK = "21323D"
MUTED = "5B6E79"
DOC_FONT = "Noto Sans"


def font_path() -> str:
    for path in [
        str(ROOT / "app/fonts/NotoSans-Regular.ttf"),
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]:
        if Path(path).exists():
            return path
    raise FileNotFoundError("Document font unavailable")


FONT_PATH = font_path()


def f(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_PATH, size)


def round_box(d: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str,
              detail: str = "", fill: str = "#FFFFFF", outline: str = "#B7CDD4",
              title_color: str = "#203B50") -> None:
    d.rounded_rectangle(box, radius=24, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = box
    cx = (x1 + x2) / 2
    cy = (y1 + y2) / 2
    if detail:
        d.text((cx, cy - 22), title, font=f(38), fill=title_color, anchor="mm")
        d.text((cx, cy + 38), detail, font=f(26), fill="#5B6E79", anchor="mm")
    else:
        d.text((cx, cy), title, font=f(37), fill=title_color, anchor="mm")


def arrow(d: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int],
          color: str = "#5D8599", width: int = 6) -> None:
    d.line((start, end), fill=color, width=width)
    theta = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 20
    spread = 0.5
    p1 = (end[0] - length * math.cos(theta - spread), end[1] - length * math.sin(theta - spread))
    p2 = (end[0] - length * math.cos(theta + spread), end[1] - length * math.sin(theta + spread))
    d.polygon([end, p1, p2], fill=color)


def make_architecture() -> None:
    im = Image.new("RGB", (2200, 1180), "#FFFFFF")
    d = ImageDraw.Draw(im)
    d.text((100, 70), "Lumori system architecture", font=f(58), fill="#203B50")
    d.text((100, 135), "Tagged deployable baseline v0.2.0-ca1 · one process, one app instance", font=f(28), fill="#5B6E79")

    round_box(d, (100, 220, 620, 385), "Browser UI", "Courses · sources · review", fill="#EDF5F5")
    round_box(d, (840, 220, 1360, 385), "FastAPI", "Auth · ownership · API", fill="#EDF3F9")
    round_box(d, (1580, 220, 2100, 385), "Persistence", "SQLite WAL + file volume", fill="#F4F7F8")
    arrow(d, (620, 302), (840, 302))
    arrow(d, (1360, 302), (1580, 302))

    d.rounded_rectangle((90, 455, 1065, 815), radius=30, fill="#F8FBFC", outline="#D1DEE4", width=3)
    d.text((125, 485), "SOURCE INGESTION", font=f(32), fill="#167E85")
    round_box(d, (125, 565, 390, 705), "Parse", "Text / MinerU", fill="#FFFFFF")
    round_box(d, (465, 565, 730, 705), "Chunk", "Page · position", fill="#FFFFFF")
    round_box(d, (805, 565, 1030, 705), "Embed", "Vector API", fill="#FFFFFF")
    arrow(d, (390, 635), (465, 635))
    arrow(d, (730, 635), (805, 635))
    d.text((126, 748), "Images also use MinerU; Docling/vision descriptions are optional", font=f(26), fill="#5B6E79")

    d.rounded_rectangle((1135, 455, 2110, 815), radius=30, fill="#F8FBFC", outline="#D1DEE4", width=3)
    d.text((1170, 485), "CONTENT WORKFLOW", font=f(32), fill="#416C88")
    round_box(d, (1170, 565, 1435, 705), "Retrieve", "Course evidence", fill="#FFFFFF")
    round_box(d, (1510, 565, 1775, 705), "Generate", "Text / Agent", fill="#FFFFFF")
    round_box(d, (1850, 565, 2075, 705), "Review", "Revise / approve", fill="#FFF9EF", outline="#E1C7A5")
    arrow(d, (1435, 635), (1510, 635))
    arrow(d, (1775, 635), (1850, 635))
    d.text((1170, 748), "Draft export: PDF / Word / MD · Approval: learner view / optional media", font=f(25), fill="#5B6E79")
    arrow(d, (1100, 385), (590, 455))
    arrow(d, (1100, 385), (1620, 455))
    arrow(d, (590, 815), (1730, 890))
    arrow(d, (1620, 815), (1730, 890))
    round_box(d, (950, 900, 1970, 1055), "Evidence and jobs persist", "Source files · chunks · revisions · calls", fill="#EDF3F9")
    d.text((100, 1105), "External APIs bill by usage; local parsing consumes host resources.", font=f(25), fill="#5B6E79")
    im.save(ASSETS / "architecture-en.png", optimize=True)


def make_data_flow() -> None:
    im = Image.new("RGB", (2200, 1080), "#FFFFFF")
    d = ImageDraw.Draw(im)
    d.text((90, 65), "Lumori core data flow", font=f(58), fill="#203B50")
    d.text((90, 126), "Course sources become a traceable index before a draft is reviewed", font=f(28), fill="#5B6E79")
    d.text((90, 225), "A · SOURCE INGESTION", font=f(34), fill="#167E85")
    top = [
        (90, 285, 450, 455, "Upload", "PDF / text / image"),
        (520, 285, 880, 455, "Parse / inspect", "Text / MinerU"),
        (950, 285, 1310, 455, "Chunk / embed", "Positions · vector API"),
        (1380, 285, 1740, 455, "Course index", "SQLite + originals"),
    ]
    for i, (x1, y1, x2, y2, title, detail) in enumerate(top):
        round_box(d, (x1, y1, x2, y2), title, detail, fill="#EDF5F5" if i == 3 else "#FFFFFF")
        if i:
            arrow(d, (top[i - 1][2], 370), (x1, 370))
    d.text((90, 570), "B · CREATION AND DELIVERY", font=f(34), fill="#416C88")
    bottom = [
        (90, 635, 450, 805, "Learning goal", "Learner level"),
        (520, 635, 880, 805, "Retrieve / rerank", "Selected course only"),
        (950, 635, 1310, 805, "Draft / check", "Lesson / Agent"),
        (1380, 635, 1740, 805, "Draft + evidence", "Job · revision"),
        (1810, 635, 2150, 805, "Review / deliver", "Draft export"),
    ]
    for i, (x1, y1, x2, y2, title, detail) in enumerate(bottom):
        round_box(d, (x1, y1, x2, y2), title, detail,
                  fill="#FFF9EF" if i == 4 else "#FFFFFF",
                  outline="#E1C7A5" if i == 4 else "#B7CDD4")
        if i:
            arrow(d, (bottom[i - 1][2], 720), (x1, 720))
    d.line(((1560, 455), (1560, 525), (700, 525)), fill="#167E85", width=6)
    arrow(d, (700, 525), (700, 635), color="#167E85")
    d.rounded_rectangle((90, 895, 2150, 1005), radius=23, fill="#F1F6FA")
    d.text((1120, 950), "Persistent evidence: ownership · job timeline · provider usage · source snapshot · revisions", font=f(29), fill="#416C88", anchor="mm")
    im.save(ASSETS / "data-flow-en.png", optimize=True)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "DFE7EB") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def keep_table_row_together(row, header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    if header:
        tr_pr.append(OxmlElement("w:tblHeader"))


def set_font(run, name: str, size: float, color: str = INK, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), name)


def add_inline(p, value: str, size: float = 9.5, color: str = INK) -> None:
    for part in re.split(r"(\*\*.*?\*\*|\*[^*]+\*|`[^`]*`)", value):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2]); set_font(run, DOC_FONT, size, color, True)
        elif part.startswith("*") and part.endswith("*"):
            run = p.add_run(part[1:-1]); set_font(run, DOC_FONT, size, color)
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1]); set_font(run, "Noto Sans Mono", max(size - 0.5, 7.5), BLUE)
        else:
            run = p.add_run(part); set_font(run, DOC_FONT, size, color)


def make_document() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    bundled_fonts = Path(
        "/Users/theoverflow/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/"
        "libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype"
    )
    FONTCONFIG.write_text(
        '<?xml version="1.0"?>\n<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">\n'
        f'<fontconfig><dir>{ROOT / "app/fonts"}</dir><dir>{bundled_fonts}</dir>'
        '<cachedir>/private/tmp/lumori-handoff-font-cache</cachedir></fontconfig>\n',
        encoding="utf-8",
    )
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.85)
    sec.right_margin = Cm(1.85)
    sec.header_distance = Cm(0.75)
    sec.footer_distance = Cm(0.75)

    normal = doc.styles["Normal"]
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    normal.font.name = DOC_FONT
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    title_style = doc.styles.add_style("Handoff Title", WD_STYLE_TYPE.PARAGRAPH)
    for name, size, before, after, col in [
        ("Handoff Title", 25, 0, 13, NAVY),
        ("Heading 1", 16, 22, 9, NAVY),
        ("Heading 2", 11.5, 14, 6, BLUE),
        ("Heading 3", 10.5, 10, 4, TEAL),
    ]:
        style = doc.styles[name]
        style.font.name = DOC_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(col)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption = doc.styles["Caption"]
    caption.font.name = DOC_FONT
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(11)
    caption.paragraph_format.keep_together = True

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("LUMORI  /  BUSINESS PLAN HANDOFF")
    set_font(run, DOC_FONT, 7.5, MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("2026.09.23  ·  ")
    set_font(run, DOC_FONT, 7.5, MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Handoff Title")
            add_inline(p, line[2:], size=25, color=NAVY)
            i += 1; continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:], size=16, color=NAVY)
            i += 1; continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:], size=11.5, color=BLUE)
            i += 1; continue
        if line.startswith("!["):
            m = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
            if not m:
                raise ValueError(f"Bad image: {line}")
            path = SOURCE.parent / m.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            with Image.open(path) as im:
                ratio = im.width / im.height
                width = min(6.7, 4.2 * ratio) if ratio < 1 else 6.7
            picture = p.add_run().add_picture(str(path), width=Inches(width))
            picture._inline.docPr.set("descr", m.group(1))
            i += 1; continue
        if line.startswith("Figure "):
            p = doc.add_paragraph(style="Caption")
            add_inline(p, line, size=8.5, color=MUTED)
            i += 1; continue
        if line.startswith("> "):
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_shading(cell, PALE)
            set_cell_border(cell, "BEDBDD")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(7)
            add_inline(p, line[2:], size=9.5, color=NAVY)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            i += 1; continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_shading(cell, "F2F6F8")
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            for n, code_line in enumerate(code):
                if n:
                    p.add_run().add_break()
                run = p.add_run(code_line)
                set_font(run, "Noto Sans Mono", 7.3, NAVY)
            i += 0; continue
        if line.startswith("|"):
            raw_rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in row):
                    raw_rows.append(row)
                i += 1
            if not raw_rows:
                continue
            ncols = max(map(len, raw_rows))
            table = doc.add_table(rows=len(raw_rows), cols=ncols)
            table.autofit = True
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ridx, row in enumerate(raw_rows):
                keep_table_row_together(table.rows[ridx], header=ridx == 0)
                for cidx in range(ncols):
                    cell = table.cell(ridx, cidx)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_shading(cell, NAVY if ridx == 0 else ("F7FAFB" if ridx % 2 == 0 else "FFFFFF"))
                    set_cell_border(cell)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    add_inline(p, row[cidx] if cidx < len(row) else "", size=7.9 if ncols >= 4 else 8.25,
                               color="FFFFFF" if ridx == 0 else INK)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.65)
            add_inline(p, line)
            i += 1; continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1; continue
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.core_properties.title = "Lumori Business Plan Handoff"
    doc.core_properties.subject = "User manual, system architecture, data flow, costs, and evidence"
    doc.core_properties.author = "Lumori"
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    ASSETS.mkdir(parents=True, exist_ok=True)
    make_architecture()
    make_data_flow()
    make_document()
