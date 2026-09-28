"""Build the formal Part-I manuscript. No model calls or application writes."""
from pathlib import Path
import hashlib, json, re, os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'docs/thesis'; FIG=BASE/'figures'; OUT=ROOT/'output'; OUT.mkdir(exist_ok=True)
# Diagram and chart assets are built separately and retained with provenance.
# Rebuilding this manuscript must not regenerate archived evidence or screenshots.
VERSION = '0.5'
fc=ROOT/'tmp/thesis/fonts.conf'
if fc.exists(): os.environ['FONTCONFIG_FILE']=str(fc)
refs=[r['citation'] for r in json.loads((BASE/'references.json').read_text())]
# Optional inspected page-map makes the contents a stable, accessible static table.
mapfile=BASE/'page_map.json'; pagemap=json.loads(mapfile.read_text()) if mapfile.exists() else {}
contents=[('List of Figures','4'),('1  Introduction','5'),('2  Literature Review','5'),('3  Requirements and System Design','7'),('4  Implementation','9'),('5  Evaluation Methodology','12'),('6  Results and Discussion','14'),('7  Limitations and Next Research Stage','19'),('References','20'),('Appendix A  AI-Assisted Content Audit','22'),('Appendix B  Evidence and Reproduction Index','24'),('Appendix C  Extended Evidence and Figure Provenance','40')]
text=(BASE/'Manuscript.md').read_text()
figure_titles=['Context position and answer accuracy','Logical system architecture','Authentication and ownership','Current difficulty allocation interface','Current citation interface','Document ingestion and source association','Chunk boundaries and provenance','Original and rewritten retrieval paths','Per-question assessment workflow','Main agent and isolated subagents','Difficulty acceptance decision','Exploration through a temporary index','Generation job and draft lifecycle','Versioned review and export','Current settings and export interface','Paired evaluation protocol','R4 held-out ranking comparison','R7 ranking across wording levels','R8A aggregate ranking comparison','R8A paired rank changes','R8B results by information sufficiency']
figure_map=pagemap.get('figures',{})
text=text.replace('{{FIGURE_LIST}}','\n\n'.join(f'{i+1}  {title} | {figure_map.get(str(i+1), "?")}' for i,title in enumerate(figure_titles)))
text=text.replace('{{CONTENTS}}','\n'.join(f'{h} | {pagemap.get(h,p)}' for h,p in contents))
text=text.replace('{{REFERENCES}}','\n\n'.join(f'[{i+1}] {v}' for i,v in enumerate(refs)))
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.3)
sec.header_distance=Cm(1.0);sec.footer_distance=Cm(1.0)
sec.different_first_page_header_footer=True
styles=doc.styles
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[name].font.name='Times New Roman';styles[name].font.color.rgb=RGBColor(0,0,0)
    rf=styles[name].element.get_or_add_rPr().rFonts
    for attr in list(rf.attrib):
        if 'Theme' in attr: del rf.attrib[attr]
    for attr in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+attr),'Times New Roman')
normal=styles['Normal'];normal.font.size=Pt(11.5)
normal.paragraph_format.line_spacing=1.10
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True
for n,size,before,after in [('Heading 1',18,20,12),('Heading 2',13,10,6),('Heading 3',11.5,9,5)]:
    s=styles[n];s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after)
    s.paragraph_format.keep_with_next=True
for sty in ['Title','Subtitle']:
    pp=styles[sty].element.find(qn('w:pPr'))
    if pp is not None:
        for b in list(pp.findall(qn('w:pBdr'))): pp.remove(b)
styles['Caption'].font.size=Pt(10);styles['Caption'].font.italic=False;styles['Caption'].font.bold=False
styles['Caption'].paragraph_format.space_after=Pt(8)
styles['Caption'].paragraph_format.line_spacing=1.05
# Running head and automatic page numbers.
hp=sec.header.paragraphs[0];hp.text='LUMORI  /  CORE TEXT GENERATION AND RETRIEVAL';hp.style='Caption'
hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor.from_string('606060')
pp=hp._p.get_or_add_pPr();border=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
for k,v in {'val':'single','sz':'4','color':'BFBFBF','space':'5'}.items():bottom.set(qn('w:'+k),v)
border.append(bottom);pp.append(border)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
# The main study uses continuous single-column academic flow. Front matter,
# bibliography and appendices retain explicit boundaries; headings stay with text.
def para(s,style=None):
    p=doc.add_paragraph(s,style)
    if style is None:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def caption(text):
    p=doc.add_paragraph(style='Caption')
    p.paragraph_format.keep_together=True
    label=re.match(r'^(Figure \d+\.)(.*)$',text)
    if label:
        p.add_run(label[1]).bold=True
        p.add_run(label[2])
    else: p.add_run(text)
    return p

def table(lines):
    data=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
    data=[r for r in data if not all(re.fullmatch(r'[-: ]+',c) for c in r)]
    t=doc.add_table(rows=0,cols=len(data[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths={3:[1.65,2.30,2.43],4:[2.2,1.40,1.39,1.39],5:[2.35,1.03,1.1,1.08,.82]}.get(len(data[0]),[6.38/len(data[0])]*len(data[0]))
    if data[0][0] in ['Condition']:widths=[2.65,1.05,.85,1.83]
    if data[0][0] in ['Task and dataset','Wording level']:widths=[3.0,1.12,1.13,1.13]
    if data[0][0]=='Strategy' and len(data[0])==5:widths=[2.32,1.06,1,1,1]
    if data[0][0]=='Input condition':widths=[2.55,1.35,1.05,1.43]
    if data[0][0]=='ID':widths=[.42,3.37,2.59]
    if data[0][0]=='Review ID':widths=[.8,.82,4.76]
    if data[0][0]=='Step':widths=[.5,5.88]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'80' if side in ['top','bottom'] else '90');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for i,values in enumerate(data):
        row=t.add_row();trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,s) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
            # Permit long evidence paths to wrap without inserting visible punctuation.
            s=s.replace('/','/\u200b') if 'docs/' in s else s
            r=p.add_run(s);r.font.size=Pt(9.5 if len(data[0])>3 or data[0][0]=='ID' else 10)
            if i==0:
                r.bold=True;shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F0');cell._tc.get_or_add_tcPr().append(shade)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)

for page_idx,chunk in enumerate(text.split('<!-- PAGE -->')):
    if page_idx==0:
        p=doc.add_paragraph('LUMORI',style='Title');p.paragraph_format.space_before=Pt(80)
        p.runs[0].font.size=Pt(34);p.runs[0].bold=True
        p=doc.add_paragraph('Course-Grounded Educational\nContent and Assessment Generation');p.paragraph_format.space_before=Pt(22);p.paragraph_format.space_after=Pt(28)
        p.runs[0].font.size=Pt(23)
        p=doc.add_paragraph('INTERIM THESIS MANUSCRIPT · PART I');p.runs[0].font.size=Pt(11);p.runs[0].bold=True
        para('Core text generation, retrieval and difficulty control')
        p=para('AI-Powered Multimodal Content & Assessment Generation for Educational Applications');p.paragraph_format.space_before=Pt(30)
        para('Final-year project · CA1 stage · AY 2026/2027')
        para(f'Version {VERSION}')
        p=para('Working thesis manuscript\nMultimodal output research will be incorporated in a subsequent stage.');p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_before=Pt(58);p.runs[0].font.size=Pt(10);p.runs[0].italic=True
        continue
    lines=chunk.strip().splitlines();i=0
    while i<len(lines):
        s=lines[i].strip();i+=1
        if not s:continue
        if s.startswith('|'):
            block=[s]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
            table(block);continue
        if s.startswith('#'):
            n=len(s)-len(s.lstrip('#'))
            p=para(s[n:].strip(),f'Heading {min(n,3)}')
            if n == 1 and not re.match(r'^[2-7]\s', s[n:].strip()):
                p.paragraph_format.page_break_before=True
                p.paragraph_format.space_before=Pt(0)
            continue
        if s.startswith('!PAIR['):
            m=re.fullmatch(r'!PAIR\[(.*)\]\((.*)\)',s)
            pair=doc.add_table(rows=1, cols=2)
            pair.autofit=False;pair.alignment=WD_TABLE_ALIGNMENT.CENTER
            pair.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,path in zip(pair.rows[0].cells,m[2].split('|')):
                cell.width=Inches(3.15)
                p=cell.paragraphs[0]
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next=True
                p.add_run().add_picture(str(BASE/path),width=Inches(2.8))
            caption(m[1]);continue
        if s.startswith('!['):
            m=re.fullmatch(r'!\[(.*)\]\((.*)\)',s);p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.line_spacing=1
            width=4.5 if 'literature-middle' in m[2] else (2.8 if '/ui-' in m[2] else 6.3)
            if 'supervisor-generated' in m[2]: width=5.9
            p.add_run().add_picture(str(BASE/m[2]),width=Inches(width))
            caption(m[1]);continue
        if (page_idx==2 or '{{FIGURE_LIST}}' in chunk or chunk.strip().startswith('# List of Figures')) and ' | ' in s:
            label,num=s.rsplit(' | ',1);p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
            from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.3),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            p.add_run(label+'\t'+num);continue
        if s.startswith('Table ') or s.startswith('Algorithm '):
            p=para(s,'Caption');p.paragraph_format.keep_with_next=True;continue
        if re.match(r'^\[\d+\]',s):
            p=para(s);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.10
            p.paragraph_format.keep_together=True
            p.paragraph_format.left_indent=Cm(.7);p.paragraph_format.first_line_indent=Cm(-.7)
            for r in p.runs:r.font.size=Pt(11.5)
            continue
        para(s)
props=doc.core_properties;props.title='Lumori: Course-Grounded Educational Content and Assessment Generation';props.subject='Interim thesis manuscript, Part I';props.author='Lumori project';props.keywords='RAG, educational generation, difficulty control, CA1';props.comments='AI-assisted manuscript. Original dates and run metadata remain in the evidence archive; no human review is asserted.'
out=OUT/f'Lumori_Thesis_Part_I_v{VERSION}.docx';doc.save(out)
embedded_paths=[]
for match in re.finditer(r'!\[.*?\]\((.*?)\)|!PAIR\[.*?\]\((.*?)\)', (BASE/'Manuscript.md').read_text()):
    embedded_paths.extend((match[1] or match[2]).split('|'))
manifest={
    'version':VERSION,
    'screenshot_provenance':'figures/v0.4/screenshot_provenance.json',
    'diagram_provenance':'figures/v0.5/vector_provenance.json',
    'notes':'Twelve original vector module diagrams; unchanged light-mode captures, licensed research chart and measured-result plots. Historical run dates remain in evidence_index_v05.json and the source archive.',
    'figures':[]
}
for rel in embedded_paths:
    p=BASE/rel
    kind=('browser_capture' if p.name.startswith('ui-') else
          'licensed_reproduction' if p.name=='literature-middle.png' else
          'original_vector_diagram' if 'v0.5' in p.parts else 'archived_result_chart')
    manifest['figures'].append({'file':str(p.relative_to(FIG)), 'kind':kind,
                               'sha256':hashlib.sha256(p.read_bytes()).hexdigest(), 'embedded':True})
(BASE/'figure_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(out)
